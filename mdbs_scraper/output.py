"""Deterministic CSV, JSON, XLSX, manifest, and project-count exports."""

from __future__ import annotations

import csv
import hashlib
import json
import logging
from pathlib import Path
from typing import Iterable, Sequence

from .errors import OptionalDependencyMissing
from .schema import ProjectRecord, STANDARD_FIELDS


LOG = logging.getLogger(__name__)

# Per-bank files written more than this far apart probably come from
# different refreshes -- combining them quietly mixes vintages.
STALE_INPUT_SECONDS = 6 * 3600


def rows(records: Iterable[ProjectRecord], *, include_raw: bool = False):
    return [record.to_row(include_raw=include_raw) for record in records]


def _fieldnames(include_raw: bool) -> list[str]:
    return STANDARD_FIELDS + (["source_fields_json"] if include_raw else [])


def write_csv(path: Path, records: Iterable[ProjectRecord], *, include_raw: bool = False) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    data = rows(records, include_raw=include_raw)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=_fieldnames(include_raw), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(data)
    return path


def write_json(path: Path, records: Iterable[ProjectRecord], *, include_raw: bool = False) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(rows(records, include_raw=include_raw), handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    return path


def write_xlsx(path: Path, records: Iterable[ProjectRecord], *, include_raw: bool = False) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise OptionalDependencyMissing(
            "XLSX output requires openpyxl. Run: python -m pip install '.[xlsx]'"
        ) from exc
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "projects"
    headers = _fieldnames(include_raw)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows(records, include_raw=include_raw):
        sheet.append([row.get(header, "") for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(12, len(header) + 2), 42)
    workbook.save(path)
    workbook.close()
    return path


WRITERS = {"csv": write_csv, "json": write_json, "xlsx": write_xlsx}


def write_records(
    path: Path, records: Iterable[ProjectRecord], output_format: str, *, include_raw: bool = False
) -> Path:
    return WRITERS[output_format](path, records, include_raw=include_raw)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def coverage_summary(
    bank_id: str, records: list[ProjectRecord], metadata: dict[str, int | str | None] | None = None
) -> dict[str, int | str | None]:
    def count(field: str) -> int:
        return sum(getattr(record, field) not in (None, "") for record in records)

    def zero_count(field: str) -> int:
        return sum(getattr(record, field) == 0 for record in records)

    summary: dict[str, int | str | None] = {
        "bank_id": bank_id,
        "status": "success",
        "error": "",
        "project_count": len(records),
        "with_commitment_year": count("commitment_year"),
        "with_loan_amount": count("loan_amount"),
        "zero_loan_amount": zero_count("loan_amount"),
        "with_total_commitment": count("total_commitment"),
        "with_total_disbursement": count("total_disbursement"),
        "with_sector": count("sector"),
        "with_sector_category": count("sector_category"),
        "with_subnational_location": sum(
            bool(record.province or record.location_text)
            or record.latitude is not None
            or record.longitude is not None
            for record in records
        ),
        "with_conditionality": count("conditionality"),
        "with_cofinancing": sum(
            bool(record.cofinancing_partners) or record.cofinancing_amount is not None
            for record in records
        ),
        "with_disbursement_timing": sum(
            bool(record.first_disbursement_date or record.last_disbursement_date)
            for record in records
        ),
    }
    if metadata:
        summary.update({
            "rows_seen": metadata.get("rows_seen"),
            "rows_retained": metadata.get("rows_retained"),
            "duplicates_removed": metadata.get("duplicates_removed"),
            "source_total": metadata.get("source_total"),
            "completeness": metadata.get("completeness", "source-total-unavailable"),
        })
    return summary


def write_summary_csv(path: Path, summaries: list[dict[str, int | str]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "bank_id", "status", "error", "project_count", "with_commitment_year", "with_loan_amount",
        "zero_loan_amount",
        "with_total_commitment", "with_total_disbursement", "with_sector", "with_sector_category",
        "with_subnational_location", "with_conditionality", "with_cofinancing",
        "with_disbursement_timing", "rows_seen", "rows_retained",
        "duplicates_removed", "source_total", "completeness",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(summaries)
    return path


def read_csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """Read one of this package's own CSV outputs: (header, rows)."""

    csv.field_size_limit(1 << 30)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def combine_csv_files(
    paths: Sequence[Path], out_path: Path, *, include_raw: bool = False
) -> list[dict[str, object]]:
    """Concatenate per-bank project CSVs into one combined CSV, in the given order.

    Every input must carry exactly this package's current column layout, so
    files written by an older schema version are refused rather than mixed
    in with shifted columns. Returns one provenance entry per input (path,
    sha256, row count, modification time) for manifest.json.
    """

    expected = _fieldnames(include_raw)
    inputs: list[dict[str, object]] = []
    all_rows: list[dict[str, str]] = []
    for path in paths:
        header, rows = read_csv_rows(path)
        if header != expected:
            missing = [name for name in expected if name not in header]
            extra = [name for name in header if name not in expected]
            raise ValueError(
                f"{path.name} does not have the current column layout "
                f"(missing: {missing or 'none'}; unexpected: {extra or 'none'}). "
                "Re-run that bank before combining."
            )
        all_rows.extend(rows)
        inputs.append({
            "path": str(path), "sha256": sha256(path), "rows": len(rows),
            "modified": path.stat().st_mtime,
        })
    times = [float(entry["modified"]) for entry in inputs]
    if times and max(times) - min(times) > STALE_INPUT_SECONDS:
        oldest = min(inputs, key=lambda entry: float(entry["modified"]))
        LOG.warning(
            "Per-bank files span %.1f hours (oldest: %s); they may come from different refreshes",
            (max(times) - min(times)) / 3600, Path(str(oldest["path"])).name,
        )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=expected)
        writer.writeheader()
        writer.writerows(all_rows)
    return inputs


def write_dict_rows(path: Path, fieldnames: Sequence[str], rows: Iterable[dict[str, object]]) -> Path:
    """Write plain dict rows (supplementary and derived datasets) as utf-8-sig CSV."""

    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fieldnames), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _plain(value) for key, value in row.items()})
    return path


def _plain(value: object) -> object:
    if value is None:
        return ""
    if hasattr(value, "is_finite") and hasattr(value, "quantize"):  # Decimal
        return format(value, "f")
    return value
